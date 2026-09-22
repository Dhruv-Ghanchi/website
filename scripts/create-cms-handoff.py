from pathlib import Path
from datetime import date
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'CMS_AI_AGENT_HANDOFF.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
doc.styles['Normal'].font.name = 'Calibri'
doc.styles['Normal'].font.size = Pt(10.5)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('236650')
section.header.paragraphs[0].text = 'GHANCHI INVESTMENTS / STRAPI IMPLEMENTATION HANDOFF'
section.footer.paragraphs[0].text = 'Future CMS specification — not a connected integration'

def heading(text, level=1):
    doc.add_heading(text, level)

def p(text):
    doc.add_paragraph(text)

def bullet(text):
    doc.add_paragraph(text, style='List Bullet')

def table(rows, headers=('Area', 'Requirements')):
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for left, right in rows:
        cells = t.add_row().cells
        cells[0].text, cells[1].text = left, right

heading('Ghanchi Investments: frontend and Strapi 5 implementation plan', 0)
p(f'Updated {date.today().isoformat()}. Working copy: C:/Users/Ghanchi/Desktop/ghanchi-investments. Branch: ghanchi-investments.')
p('The verified Kora replica remains on main at commit d4350cf in https://github.com/Dhruv-Ghanchi/website. Never rebrand that backup. This document replaces the historical Supabase/custom-admin architecture with Strapi 5. No CMS, hosting, database, media service or delivery provider has been provisioned.')
heading('1. Scope, current state and non-negotiable requirements')
for text in [
    'Implemented frontend: Ghanchi canonical routes, nine service records, About subpages, online-service links, dated newsletters, curated archive article adaptations, contact details, forms and legal drafts. Content remains file-backed. Strapi schemas, API fetching, preview, revalidation and editor workflows are NOT implemented.',
    'Preserve Manrope typography, mint/cream/ink tokens, rounded panels, pill buttons, image hovers, reveal springs, reduced-motion behavior, scroll pinning, comparison scene, process accordion/dialog, advisor dialog, testimonial expansion and footer circle/scale/hover interactions.',
    'Only pricing removal was approved. Other panels must be repurposed, not removed. Obtain owner approval before future removal of any component, feature or section.',
    'The owner approved 1,200+ clients, 15+ years, 12+ awards and a static 5.0 Google rating from the existing website. These are editorial values, not independently audited measurements. Do not invent review counts or claim live synchronization.',
    'Routine editorial changes must be possible in Strapi without frontend code changes. Entirely new section types, integrations or executable behavior still require development.',
    'Keep this generator and generated DOCX aligned. Regenerate with python scripts/create-cms-handoff.py after specification changes.'
]: bullet(text)
heading('2. Source audit, provenance and migration decisions')
table([
    ('Business source', 'https://ghanchiinvest.com/ and /about-us/: established 2009; Chandrakant B. Ghanchi; 1,200+ clients; goal context, risk appetite, cash flows, implementation assistance and annual review.'),
    ('Verified contact', '+91 9820926446; +91 7977061717; info@ghanchiinvest.com; chandrakant@ghanchiinvest.com; Shop no. 27, Sector 11, Balaji Bhavan, CBD Belapur, Navi Mumbai, Maharashtra 400614.'),
    ('Client locations', 'India, UAE and USA are supported by source content/testimonials. Do not add countries from assumptions. Employer affiliations are not corporate customer endorsements.'),
    ('Testimonials', 'Use actual attributed text from /testimonials/ and the homepage. Preserve spelling or record approved editorial corrections. Portraits are absent: use initials, not unrelated stock faces. All eleven testimonials identified across the source homepage and testimonial page retain source URLs; long quotes may use attributed excerpts.'),
    ('Awards and certificates', 'Imported eight award photos and nineteen certificate scans. Store neutral archive labels until issuer, date, title and validity are individually transcribed and approved. Historical certificates must not imply current regulatory authorization.'),
    ('Newsletter archive', 'Nineteen entries, May 2020 through November 2021. Latest means newest listed issue, not a current upload. November 2020 and June 2020 links are malformed: keep records visible with unavailable state; obtain correct URLs from owner, do not guess.'),
    ('Articles', 'Two legitimate March 2021 articles are currently condensed, labelled adaptations. Preserve original dates, source URLs and adaptation notes. They are not newly authored by the founder. Additional legitimate archive articles remain editorial migration work.'),
    ('Excluded material', 'Unrelated casino/gambling posts appear on the live site. Do not copy lists, related posts, HTML scripts or bulk WordPress exports without review. Existing-site security remediation is a launch prerequisite with the current maintainer.'),
    ('Imagery', 'Founder, awards and certificate assets come from Ghanchi. Other retained Kora-derived photographs are illustrative, not Ghanchi staff or customer identities. Confirm reuse rights and choose owner-approved replacements before launch.'),
    ('Financial copy', 'Current service summaries are conservative adaptations. Validate current product, tax and regulatory wording, registered designations and required registration numbers with the owner. Never add guaranteed returns or regulatory claims from inference.'),
    ('Online services', 'Links are copied from header/footer sources, with the Niva Bupa footer destination used instead of the old MaxBupa route. Provider uptime and authenticated journeys are not certified. No live login or payment transaction is part of testing.')
])
heading('3. Kora component-to-business mapping')
table([
    ('Shared shell', 'Six primary items: Home, About Us, Services, Online Services, Blog, Contact Us. Desktop disclosure buttons support keyboard input and Escape; mobile links close navigation. Keep footer wordmark, circle reveal, hover transition, social links and clipboard status.'),
    ('Hero', 'Preserve measured-height scroll timeline, pinned background scale/blur, word reveal and foreground inert state. Use Ghanchi heading/actions. Retain avatar row with initials, client count and honest rating label. Retain marquee with approved client locations.'),
    ('Featured hero card', 'Newsletter relationship, not case study. Resolve latest published issue by issueMonth descending. Show original month/year and archive label. Never synthesize a newer issue. Empty/unavailable state must remain visible and useful.'),
    ('Statement and comparison', 'Treat hero, statement, scroll-space and comparison as one indivisible scene. Editable statement and before/after copy; preserve transforms and breakpoints. Do not promise lifetime corpus or certainty of outcomes.'),
    ('Services and chart', 'Nine service cards with original number layout, hover image, glass overlay and CTA. The animated chart illustrates the planning journey; it must not imply measured investment returns. Editable labels and disclosure; no invented values.'),
    ('Process and principles', 'Understand, Assess, Plan, Review. Keep spring accordion, play-button dialog and operating-principle cards. Video is unconfigured: show explicit text fallback plus real YouTube channel link, not a fictitious duration.'),
    ('Advisor and former hiring panel', 'One verified founder profile with real photo and biography. Preserve expandable modal. Former hiring panel becomes awards archive panel using genuine photo and CTA.'),
    ('Testimonials and counters', 'Actual client quotes, initials if no approved portrait, source metadata, expandable cards. Counters use owner-selected values. Founder callout remains. Google feed remains unconnected.'),
    ('Former featured case', 'Preserve narrative, facts, large statistics, service chips, image and button; present Ghanchi client community. No fictional company performance or investment result.'),
    ('Pricing', 'Removed under explicit approval. Do not create premiums, advisory packages or subscription prices without owner-approved product information.'),
    ('FAQ, blog and contact', 'Keep tabbed FAQ, accordion semantics, image-led article cards and full contact panel. Adapt all business labels. Preserve consent, honeypot, prehydration safety, validation, loading and honest delivery failures.')
])
heading('4. Canonical route inventory and legacy behavior')
for text in [
    '/; /about-us; /about-us/awards; /about-us/certificates; /about-us/our-clients; /about-us/testimonials.',
    '/services and /services/[slug] for all nine services, using lowercase hyphenated slugs defined in content.ts.',
    '/online-services; /newsletters; /blog; /blog/[slug]; /contact-us; /privacy-policy; /terms-of-service; /disclaimer.',
    'Legacy /about -> /about-us; /contact -> /contact-us with encoded service query retained; /insights -> /blog; /cases -> /about-us/our-clients. Fictional case detail slugs return 404 instead of misrepresenting Ghanchi clients.',
    'Legacy root /awards, /certificates, /our-clients and /testimonials redirect to About subpages. /blog-post redirects to /blog; known root service and article slugs redirect to canonical content. Unknown slugs return 404.',
    'SITE_URL controls metadataBase and sitemap origin. Without SITE_URL, robots disallows indexing for preview. Do not enable production crawling until legal/content/integration approval. Add canonical metadata for any remaining per-record routes and verify host redirects at deployment.'
]: bullet(text)
heading('5. Architecture: typed server-side content delivery')
p('Strapi runs as a separate application and administrator interface. Next.js remains the public frontend. Use Strapi 5 REST API with explicit fields, bounded pagination and relation population. GraphQL is optional, not needed by default. Strapi 5 uses documentId and flattened response fields; do not copy Strapi 4 data.attributes adapters.')
for text in [
    'Define ContentRepository methods: getSiteSettings, getNavigation, getHomepage, getPageByPath, listServices, getServiceBySlug, listArticles, getArticleBySlug, listTestimonials, listNewsletters, getFooter and getContactSettings. Provide local and Strapi adapters returning identical domain types.',
    'Introduce a server-only module for STRAPI_URL and STRAPI_READ_TOKEN. Fetch published status by default. Never expose tokens via NEXT_PUBLIC variables or import the repository into client modules.',
    'Refactor current client components to accept serialized props from Server Components. Today they import content.ts directly, which bundles collections. Do not claim the CMS is connected until those imports have been replaced and tested.',
    'Validate API responses at the adapter boundary: required fields, unique slugs, date formats, enum values, URL schemes/hosts, media dimensions, references and section limits. Fail closed on unknown __component IDs. Render approved safe empty states for optional missing records; do not silently substitute fictional content.',
    'Fetch shared settings/navigation once per render with request memoization. Use Next.js cache tags for site, navigation, homepage, service:slug, article:slug and galleries. Define the cache APIs for the installed Next.js release rather than assuming older revalidateTag signatures.',
    'If Strapi is down, serve the last valid published cache where available and surface operational alerts. Never expose draft content as a fallback. Explicit development mode may use local seed data; production must not silently switch adapters.'
]: bullet(text)
heading('6. Strapi single types: complete field model')
table([
    ('Site Settings', 'siteName required string(100); logo media image; logoAlt string(180); wordmark string(40); locale enum en-IN; canonicalBaseUrl HTTPS; defaultSEO component; socialLinks repeatable link; legalDisclaimer text(2000); approvedTheme enum mint-cream; favicon media. Theme tokens are owner-only bounded settings, not arbitrary CSS.'),
    ('Navigation', 'items repeatable navigation-item, max 8; each has stable key, label(40), internal path or external link, enabled boolean, children max 16 and one nested level only. Order is editorial. Include mobile CTA, header contact CTA, optional advisor relationship, skip-link label. Validate links and unique keys before publishing.'),
    ('Homepage', 'title; SEO; required intro-scene component containing hero, statement and comparison; body dynamic zone allowlisted sections in approved order; relation picks and enabled states. Intro scene cannot be split/reordered. Required contact path and no duplicate H1. Initial section order matches frontend mapping.'),
    ('Contact Settings', 'heading lines, introduction, benefit list, testimonial relation, phone/email/address components, form labels/help text, submit/loading/success/error labels, service picklist relation, consent copy/link, privacy warning, statistic relations. Required fields and security validation remain code-defined; editors cannot disable consent or honeypot.'),
    ('Footer Settings', 'newsletter heading/description, consent text and privacy link, contact statement, advisor relation, address, phone/email links, social link relations, legal page relations, navigation relation, wordmark, copyright business name and disclosure. Email-copy status labels editable; clipboard behavior stays in code.')
])
heading('7. Strapi collection types and relationships')
table([
    ('Page', 'title string(120); slug UID unique; path unique internal path; template enum standard/about/gallery/clients; SEO; allowed sections dynamic zone; navigationTitle; enabled; sourceUrl; editorialReviewDate. Dynamic pages cannot shadow reserved API or legal routes.'),
    ('Service', 'title string(80); slug UID unique; shortDescription(140); longDescription(1400); body rich-text blocks; icon enum of existing lucide mappings; image media + alt/focal point; benefits repeatable string(120), 1..12; general testimonial relation (clearly not product-specific unless approved); relatedServices many-to-many; CTA link; sortOrder integer; sourceUrl; SEO.'),
    ('Advisor', 'name(100), role(100), slug UID, biography text(3000), headshot media with identity approval, social links, sortOrder, profileEnabled. No stock image may be attributed to an actual advisor. One founder record initially.'),
    ('Testimonial', 'name(100), role/affiliation(180), location optional, quote text(2500), approvedExcerpt optional, portrait optional, sourceUrl required, consent/ownerApproval metadata, approved boolean, sortOrder. Preserve full quote; distinguish employment affiliation from endorsement.'),
    ('Award and Certificate', 'Separate collections: title(160), slug UID, image required, alt required, issuer optional, issueDate optional, expiryDate optional, description(1500), sourceUrl, archive flag, verificationState enum unreviewed/owner-approved. Do not invent missing issuer, dates or licence status. Show neutral labels until transcribed.'),
    ('Client Reference', 'displayName, kind enum individual-affiliation/approved-corporate, description, location relation, logo optional with usage approval, consent evidence private, sortOrder, sourceUrl. Corporate-customer designation requires explicit approval, not inference from Our Clients logos.'),
    ('Article', 'title(180), slug UID, excerpt(500), blocks, category relation, authorLabel or verified advisor relation, cover image/alt, originalPublicationDate, sourceUrl, adaptationNote, reviewDate, financialReviewStatus, SEO, featured boolean. Preserve historical dates; exclude spam. Original admin byline must not silently become founder authorship.'),
    ('Category', 'name(60), slug UID unique, description(300), sortOrder. Empty category state supported.'),
    ('Newsletter', 'title(140), slug UID, issueMonth YYYY-MM, cover image optional, providerUrl HTTPS allowlisted or PDF media, availability enum available/unavailable, unavailableReason, originalSourceUrl, sortOrder. Exactly one destination for available records. Latest relation resolved by issueMonth, never arbitrary array order.'),
    ('Online Service', 'title(80), slug UID, description(300), destination safe link, providerName, type enum account/renewal/newsletter/app, opensNewTab boolean for external only, logo optional, enabled, sortOrder, verifiedAt optional, sourceUrl. No customer credentials stored in CMS.'),
    ('Statistic', 'key UID, numericValue decimal, prefix/suffix(12), decimals integer 0..2, label(180), sourceUrl, asOfDate optional, ownerApproved boolean, format enum number/rating/year. Counter formatting and accessible labels derive from the same record.'),
    ('FAQ', 'question(180), answer blocks/text(2500), category enum General/Planning/Protection/Online Services, sortOrder, published boolean through Draft & Publish. Stable IDs drive ARIA connections.'),
    ('Legal Page', 'title, slug enum privacy-policy/terms-of-service/disclaimer, body blocks, effectiveDate optional, approvalStatus, approvedBy private, SEO noindex until approved. Publish gate must prevent draft legal text becoming approved merely by changing a date.'),
    ('Client Location', 'name(60), countryCode optional, sourceUrl, approved boolean, sortOrder. Initially India, UAE, USA only. No personal addresses.'),
    ('Media metadata', 'Use Media Library files plus alt/caption, focalX/Y numbers 0..100, source/rights record, identity consent flag when depicting a named person, original filename, width/height. Keep source provenance separate from public caption.')
])
heading('8. Reusable components and validation boundaries')
table([
    ('seo.metadata', 'metaTitle <=60 recommended; metaDescription <=160 recommended; canonical internal path only; OG image; noindex boolean; social title/description overrides. Render plain text and generated metadata, not scripts supplied by editors.'),
    ('shared.link', 'label required; kind enum internal/external/email/phone/anchor; destination; optional relation; newTab for external; ariaLabel optional. Validate each kind: internal must begin with a single slash, reject backslashes/control chars and protocol-relative URLs; HTTPS external; constrained mailto/tel; no javascript/data/file URLs.'),
    ('shared.media', 'Image or video media relation; alt required except decorative; caption; focal point; approved aspect preset; video poster; captions file; real duration. Validate image MIME and actual bytes, file size/dimensions, malware scan documents. SVG must be sanitized or disallowed.'),
    ('shared.motion', 'Preset enum none/reveal/accordion/hero-scene/footer-scene. Optional delay 0..0.6 seconds, ticker speed bounded and layout preset allowlisted. Reduced motion always overrides settings. No arbitrary transforms, JS, CSS, easing code or timeline scripts.'),
    ('shared.rich-text', 'Typed headings, paragraphs, lists, emphasis and safe links. Validate heading hierarchy. Sanitize at input and render using React elements; never pass untrusted HTML to dangerouslySetInnerHTML. Disable arbitrary embeds and iframes.'),
    ('shared.provenance', 'sourceUrl, sourceDate optional, reviewedAt, reviewedBy private, approvalState, notes private. Financial claims and identity-bearing media require owner review.'),
    ('shared.action', 'Enum navigate/open-profile/open-process/submit-contact/subscribe/copy-email. Editors choose labels and allowed targets; behavior and permission checks remain in code.')
])
heading('9. Allowlisted section registry and minute-level editor controls')
table([
    ('intro-scene / hero', 'Heading text and approved line-break hints; subheading; description; primary/secondary CTA; image/video/poster; decorative alt; overlay 0..0.65; trust name initials/approved portraits; statistic and rating relation; rating source note; location relations and order; newsletter mode latest/manual; badge, footer labels; scene remains code-defined.'),
    ('intro-scene / statement-comparison', 'Statement; before/after labels and heading; 2..6 comparison bullets; icon enum check/cross; light/dark logo variants. Recalculate word timing safely for editable word counts. Bound title length and line count so content does not escape fixed panels.'),
    ('services-showcase', 'Section title, intro, highlighted phrase, footnote; ordered 1..9 service relations; number style; benefits label; CTA labels; image overlay mode vision/testimonial; chart labels and disclosure. Chart interpretation is not editor-programmable.'),
    ('process-accordion', 'Heading/subheading; 2..6 phases with title/icon/body; active default; process poster/video/captions; play/close labels; fallback heading/copy; channel link; operating principles 1..4 with icon/title/body.'),
    ('advisor-showcase', 'Section heading, biography label, ordered advisor relations, profile link label, dialog close label; awards-panel image/title/description/CTA. Actual staffing controls number of rows; do not populate fictitious people to match original counts.'),
    ('testimonials', 'Featured relation; background; excerpt/full quote display; ordered expandable relations; rating relation/source note; statistic relations; founder callout heading/body/advisor/CTA. Optional portrait falls back to initials. Not all testimonials imply a five-star score.'),
    ('client-feature', 'Eyebrow/date label; title; three fact label/value pairs; two statistic relations; image/alt; services relations/chip limit; testimonial optional; CTA. Display data from one canonical record, not duplicated facts.'),
    ('faq-accordion', 'Heading; categories/order; question relations/order; initial open item; advisor help card and link. DOM IDs are generated, not editorial values.'),
    ('article-list', 'Heading, description, CTA, category filter, article relation/manual or latest mode, item limit, cover crop; empty-state copy. Show original dates and adaptation status.'),
    ('gallery', 'Title/intro, archive note, award/certificate relations, order, caption visibility, alt, image fit enum contain/cover, full-size link label. Avoid captions inventing award details.'),
    ('online-services / newsletters', 'Intro/disclosure, ordered relation list, external-link label, availability states, app-link labels and icons. Avoid rendering malformed destinations as clickable links.'),
    ('contact-cta', 'Contact settings relation; heading, intro, consent labels and link; supported statistic relations; optional contact-details panel. Delivery secrets never editable as content.'),
    ('page section envelope', 'Stable key; enabled boolean; template enum; theme variant mint/cream/paper/ink; bounded spacing preset; safe motion preset. Enable, disable and reorder only supported sections. Preview length limits and refuse unknown component IDs.')
])
p('Implementation pattern: switch on a validated discriminated union of section.__component; pass typed data to existing React components. Unknown types block publishing and produce a safe development error, not dynamic imports from user-provided strings. The entire intro-scene stays required and indivisible. Section removal beyond pricing requires explicit owner authorization even if the future UI supports enabled=false.')
heading('10. Native Strapi features versus custom work and paid plans')
table([
    ('Native foundation', 'Content-type Builder (development), Content Manager, components/dynamic zones, relations, Media Library, REST API, Draft & Publish and administrator roles/permissions. Enable Draft & Publish for editorial records. Public write/create/delete permissions remain disabled.'),
    ('Preview', 'Basic Preview is free, but requires configuring Strapi preview URLs and implementing the Next.js draft-mode endpoint and draft fetching. Live Preview is a CMS Growth/Enterprise feature. Neither frontend preview mode nor Strapi preview config is implemented here.'),
    ('Content History', 'CMS Growth or Enterprise, per official docs. It stores versions and restores into draft; it is not a replacement for database/media backups.'),
    ('Releases and scheduling', 'CMS Growth/Enterprise. Groups publish/unpublish actions and supports scheduling with timezone. Draft & Publish must be enabled. Confirm plan entitlement before committing to launch workflows.'),
    ('Review Workflows', 'CMS Enterprise. Multi-stage review pipelines and stage permissions. Do not promise this on the free edition. On a lower plan, use manual editor/publisher roles and a custom review-status field with an enforced publication gate.'),
    ('Custom engineering', 'Typed adapters, media safety, frontend section registry, relationship validation, owner-approval gates, preview security, cache invalidation, URL checks, responsive constraints, form delivery, rate limiting, Google rating synchronization and full integration tests.'),
    ('Hosting versus CMS license', 'Strapi Cloud hosting and CMS feature licensing are separate choices. Confirm exact current subscriptions, quotas and plan terms before purchase. No hosting provider or plan is chosen by this specification.')
])
for url in ['https://docs.strapi.io/cms/features/draft-and-publish', 'https://docs.strapi.io/cms/features/preview', 'https://docs.strapi.io/cms/features/content-history', 'https://docs.strapi.io/cms/features/releases', 'https://docs.strapi.io/cms/features/review-workflows']:
    p(url)
heading('11. Roles, approval and publishing')
for text in [
    'Owner/Super Admin: manages configuration, administrator invitations, allowed presentation settings, provider setup and final business/legal approval. Do not use the super-admin account for routine editing.',
    'Editor: create/update drafts for permitted content types, upload approved media, preview. No publishing, secrets, plugin installation, user management or schema modification.',
    'Publisher: review provenance, financial language, link safety and presentation; publish/unpublish approved content. Paid workflow stage permissions are configured only when entitled.',
    'Before publish: validate all references resolve to published records, required sections exist, slugs are unique, internal links resolve, external hosts are allowed, media has alt/rights metadata, and no claim is awaiting approval. Enforce server-side, not merely as a UI checklist.',
    'Publishing service/article changes invalidates detail, listing, homepage and navigation caches where referenced. Unpublishing must remove the record from listings and invalidate its detail route; define 404 or approved redirect behavior.',
    'Versioned rollback: use entitled Content History or a separately designed revision/export process. Restore drafts, preview, approve and republish. Back up database and media together and test restores in staging.'
]: bullet(text)
heading('12. Preview and revalidation security')
for text in [
    'Strapi generates a preview URL pointing at an authenticated Next.js /api/draft endpoint. Use a signed short-lived payload containing documentId, content type, allowed path, timestamp and nonce. Verify signature with constant-time comparison, expiration and replay controls. Never accept arbitrary external return URLs.',
    'Enable draftMode only after verification. Fetch draft status with a server-only scoped preview token, use no-store, set noindex and do not share draft responses through public caches. Provide a visible exit-preview route that disables draft mode.',
    'Signed revalidation is custom work: standard configurable Strapi webhook headers are not automatically per-payload HMAC signatures. Implement an explicit Strapi signing middleware/relay for raw body plus timestamp. Do not describe a plain static header as signed payload verification.',
    'Webhook receiver verifies signature, timestamp, bounded body size, event allowlist, content type and document ID. Deduplicate delivery IDs, derive affected cache tags server-side, reject arbitrary paths/tags. Use short timeouts, retries/backoff and sanitized operational logs.',
    'Test forged signatures, replay, expired preview links, draft token exposure, open redirects, concurrent publishing, unpublish, slug changes and cache leakage. Secrets are stored only in deployment environment, never CMS content or browser bundles.'
]: bullet(text)
heading('13. Contact, newsletter and rating integrations')
p('Current contact payload: name 2..100; email <=254 and valid syntax; phone 7..25 characters with at least seven digits; services array of unique known IDs, at least one; goal optional <=200; message 10..3000; consent exactly true; website honeypot empty. Newsletter: email and separate explicit consent, empty honeypot. Server enforces same-origin JSON, bounded request body and HTTPS configured destination; upstream calls use timeout and reject redirects. Missing configuration returns 503, not success.')
for text in [
    'CONTACT_WEBHOOK_URL and NEWSLETTER_WEBHOOK_URL remain unconfigured. Owner chooses delivery provider and recipients. Add distributed rate limiting, anti-abuse monitoring and provider authentication before public launch. Honeypot alone is not production spam protection.',
    'Log request IDs and outcome categories only, not message bodies, email addresses, phone numbers, tokens or webhook URLs. Agree retention, deletion and processor terms; update privacy policy accordingly.',
    'Do not store form enquiries in public Strapi collections. If a private CRM/inbox is desired, design separate least-privilege access and retention rules. Never expose credentials or investment account numbers in content APIs.',
    'Google rating: current static owner-approved 5.0 only. Future integration needs verified business/place identity, permitted API access, billing if applicable, caching and attribution compliant with provider policy. Show fetchedAt and unavailable/stale state. Never scrape customer accounts or fabricate review counts.',
    'External renewal and account links are outbound navigation only. Test link syntax and approved host, not real sign-in, premium payment or transactions.'
]: bullet(text)
heading('14. Deployment configuration and operational checklist')
table([
    ('Next.js environment', 'SITE_URL public canonical origin; server-only STRAPI_URL, STRAPI_READ_TOKEN, STRAPI_PREVIEW_TOKEN, PREVIEW_SIGNING_SECRET, REVALIDATION_SIGNING_SECRET; CONTACT_WEBHOOK_URL; NEWSLETTER_WEBHOOK_URL. Future CMS variables are specification only, not read by the current frontend.'),
    ('Strapi environment', 'Database connection, app keys, token salts, admin JWT secret and supported production settings generated securely. FRONTEND_URL and signing secrets for preview/webhook integration. Media provider credentials server-side. Never commit .env files.'),
    ('Infrastructure choices', 'Owner selects Strapi hosting, database engine/provider, persistent media/object storage, region, backup retention, SMTP/delivery, domain, licensing and monitoring. No provisioning/deployment authorized by this document.'),
    ('Environment separation', 'Independent development, staging and production data/tokens. No production API writes from tests. Verify CORS only permits approved frontend/admin origins; protect admin access; TLS; restricted database network; least-privilege tokens.'),
    ('Media durability', 'Do not use ephemeral deployment filesystem for permanent uploads. Back up media with database references, verify restoration and avoid deleting referenced assets. Use provider transformations only through safe mappings.'),
    ('Release gate', 'Run all local tests and content checks, owner visual/content sign-off, legal approval, source-site security cleanup, external-link review, delivery sandbox test, rate limiting, backup restore rehearsal, robots/canonical check and rollback rehearsal before DNS cutover.')
])
heading('15. Step-by-step implementation sequence')
for index, text in enumerate([
    'Capture approved Ghanchi frontend screenshots and motion checkpoints at 1440/1200/810/390. Preserve backup branch. Inventory all hardcoded UI labels including validation, dialogs, empty states and disclosures.',
    'Choose infrastructure and plan with owner. Pin a supported Strapi 5 release at least seven days old and compatible Node/database versions. Create isolated Strapi repository/application and secure environment variables.',
    'Create components, collections and single types in development; commit schema files. Configure Draft & Publish and roles. Seed only validated local records and approved media with provenance. Use deterministic source keys for idempotent imports.',
    'Implement publishing validation and claim/identity approval gates. Test role separation and public API permissions before exposing any data.',
    'Build server-only REST repository and runtime adapters. Refactor client renderers to typed props, beginning with settings/navigation, then hero, service/process, testimonials and inner pages. Verify design after each area.',
    'Implement safe section registry. Preserve intro grouping. Enforce limits and safe preset enums. Add editor descriptions/examples for every field and instructions for missing media and unavailable links.',
    'Implement preview with signed links, draft isolation and exit action. Test ordinary visitors cannot access drafts.',
    'Implement signed revalidation bridge and receiver, dependent tag invalidation, retry and unpublish behavior. Test stale cache and Strapi outage handling.',
    'Connect delivery provider in sandbox, add rate limits and monitoring, update privacy/legal text. Connect Google only if approved setup is available; otherwise retain the honest static label.',
    'Train editor: edit a heading, change an image/alt, reorder supported sections, add navigation child, publish service/article/newsletter, preview, approve, unpublish and restore. Verify all without code changes.',
    'Run acceptance suite, review accessibility and screenshots, resolve launch checklist. Deploy staging, obtain approval, back up, then separately authorize production deployment and DNS changes.'
], 1):
    heading(f'Step {index}', 2)
    p(text)
heading('16. Verification and acceptance tests')
for text in [
    'node scripts/complete-content.mjs --validate: IDs, required nine services, sourced content, newsletter chronology, approved locations, safe links and local asset existence. Import script now refuses non-validation mode; never restore the Kora live overwrite into this worktree.',
    'npm run typecheck; npm run build; npm test. Playwright uses isolated localhost:3100 with reuseExistingServer=false so it cannot accidentally test the Kora backup on port 3000.',
    'All canonical routes: one H1, functioning skip target, Ghanchi metadata, no broken images, no horizontal overflow, no Kora business claims or unrelated spam. Legacy destinations and unknown 404s tested.',
    '1440, 1200, 810 and 390 widths: hero pin/blur/statement, comparison, all nine service cards, process transitions/dialog, founder modal, awards/gallery, testimonial expansion, FAQ tabs, newsletter card, footer scale reversal and navigation/social hover. Reduced motion remains readable.',
    'Keyboard: dropdown disclosure Enter/Space, links reachable by Tab, Escape closes and restores focus; mobile parent and child navigation, scrollable long submenu, gallery links, dialogs, FAQ arrows and consent controls.',
    'Forms: disabled until hydration, POST not GET, no personal data in URLs, required fields, phone bounds, known service IDs, consent, honeypot, same-origin/type/size failures, network errors, loading, retained values after failure, retry and honest success. Mock every transport; no real messages.',
    'CMS acceptance: owner can edit every inventoried field, reorder/toggle supported sections, update navigation/footer and replace media without code; unknown blocks and unsafe URLs rejected; unpublished dependencies blocked; drafts remain private; publish/unpublish revalidates correct pages.',
    'Final visual sign-off is not implied by passing automated tests. Human review must compare full scrolled pages and animation checkpoints against the approved baseline. Long copy and increased service count necessarily change page length.'
]: bullet(text)
heading('17. Remaining owner decisions and honest implementation status')
for text in [
    'Strapi host, database, storage, paid plan and editor accounts are undecided and unprovisioned.',
    'Strapi schemas, content adapters, preview and signed webhooks are future work. Current content is local and not editable through Strapi yet.',
    'Google synchronization, form delivery, production abuse controls, current legal/regulatory wording and final licensing/rights approval are outstanding launch dependencies.',
    'Newsletter archive needs correct November 2020 and June 2020 URLs and any newer editions the owner wishes to provide.',
    'Awards/certificate titles, dates, issuers and expiry should be transcribed/approved; neutral archive captions remain until then. Additional authentic testimonials and archive articles can be migrated through the reviewed process.',
    'Hero flower and other illustrative photographs retain the original visual language; choose replacements only with owner approval. Current founder photo is genuine but composited source artwork, not a newly commissioned headshot.',
    'Clean up the existing WordPress site and investigate unauthorized content before relying on it for future automated imports. Do not bulk-migrate its current blog feed.'
]: bullet(text)
doc.save(OUTPUT)
with ZipFile(OUTPUT) as archive:
    ElementTree.fromstring(archive.read('word/document.xml'))
loaded = Document(OUTPUT)
text = '\n'.join(p.text for p in loaded.paragraphs)
for required in ['Strapi 5', 'Implementation', 'Preview', 'Verification', 'Remaining owner decisions']:
    if required.lower() not in text.lower():
        raise RuntimeError(f'Missing handoff section: {required}')
print(f'Generated and validated {OUTPUT.name}: {len(loaded.paragraphs)} paragraphs, {len(loaded.tables)} tables.')
