from pathlib import Path
from datetime import date
from zipfile import ZipFile
from xml.etree import ElementTree
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / 'GHANCHI_INVESTMENTS_TRANSFORMATION_PLAN.docx'
doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.75)
normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(7)
for name in ['Heading 1', 'Heading 2', 'Heading 3']:
    doc.styles[name].font.color.rgb = RGBColor.from_string('236650')
header = section.header.paragraphs[0]
header.text = 'GHANCHI INVESTMENTS / KORA-TO-GHANCHI TRANSFORMATION PLAN'
header.style = 'Caption'
footer = section.footer.paragraphs[0]
footer.text = 'Planning document — implementation requires owner approval'

def heading(text, level=1):
    doc.add_heading(text, level)

def p(text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    if bold_prefix and text.startswith(bold_prefix):
        paragraph.add_run(bold_prefix).bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)

def bullet(text, level=0):
    style = 'List Bullet' if level == 0 else 'List Bullet 2'
    doc.add_paragraph(text, style=style)

def numbered(text):
    doc.add_paragraph(text, style='List Number')

def table(rows, headers=('Kora component', 'Ghanchi Investments use')):
    t = doc.add_table(rows=1, cols=2)
    t.style = 'Light Shading Accent 1'
    for cell, text in zip(t.rows[0].cells, headers):
        cell.text = text
    for left, right in rows:
        cells = t.add_row().cells
        cells[0].text = left
        cells[1].text = right

heading('Kora Website to Ghanchi Investments', 0)
title = doc.paragraphs[-1]
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle = doc.add_paragraph('Complete content, component, route, interaction, and future-CMS transformation plan')
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
p(f'Prepared {date.today().isoformat()} for the isolated Ghanchi Investments worktree at C:/Users/Ghanchi/Desktop/ghanchi-investments.')
p('Approval status: planning document only. Further implementation should pause until the owner reviews and approves or changes the mappings in this document. The verified Kora backup on the main branch must remain untouched.')

heading('1. Transformation objective')
p('The final website will be Ghanchi Investments presented through Kora’s visual language and interaction quality. It will not be a generic financial-services redesign, nor a superficial logo replacement. Business content, information architecture, routes, services, customer journeys, claims, people, contact details, articles, and legal disclosures will become Ghanchi-specific while the Kora design system and motion behavior remain intact.')
heading('Visual and behavioral elements that must remain')
for item in [
    'Manrope typography, existing weight treatment, display scale, letter spacing, and responsive type behavior.',
    'Mint, cream, paper, white, muted-gray, and dark-ink color system.',
    'Large rounded panels, rounded cards, pill buttons, glass overlays, and existing container widths.',
    'Fixed floating header, desktop navigation pill, dark contact CTA, and animated mobile navigation.',
    'Hero entrance scale, pinned background, scroll-linked scale and blur, word-by-word reveal, and foreground movement.',
    'Animated marquee, sticky statement scene, and before/after comparison scene.',
    'Service card proportions, numbering, split layout, image hover, tags, overlay card, and reveal animations.',
    'Process accordion spring transitions, video/fallback dialog, advisor profile dialog, testimonial expansion, and FAQ tabs.',
    'Animated counters, editorial cards, contact-panel presentation, footer circle reveal, footer scale, and layered link hovers.',
    'Keyboard focus, semantic controls, Escape behavior, mobile responsiveness, and reduced-motion fallbacks.'
]: bullet(item)

heading('2. Section decision framework')
p('Every Kora section must be evaluated using the following order. A section is not removed merely because its original consulting content does not apply.')
numbered('Keep the component unchanged visually when its interaction and layout work for Ghanchi Investments.')
numbered('Repurpose the component when the original business meaning is unsuitable but the component can communicate genuine Ghanchi information.')
numbered('Remove a section only when it cannot honestly represent the business and the owner explicitly approves removal.')
numbered('Never fill an unsuitable component with invented prices, returns, customers, awards, staff, reviews, locations, or regulatory claims merely to preserve its original content density.')
p('Current approved removal: the Kora pricing section only. All other major Kora panels are retained and repurposed. Any future removal requires owner approval.')
table([
    ('Client-logo marquee', 'Verified client locations.'),
    ('Hero case-study card', 'Latest verified newsletter in the archive.'),
    ('Revenue comparison chart', 'Illustrative financial-planning journey, clearly labelled as not a return forecast.'),
    ('Growth before/after scene', 'Before and after structured, goal-based financial planning.'),
    ('Team component', 'Chandrakant B. Ghanchi and verified staff or advisors only.'),
    ('Hiring panel', 'Awards and recognition archive.'),
    ('Featured case study', 'Client community, business reach, service facts, and credibility without performance claims.'),
    ('Growth process', 'Understand, Assess, Plan, Review advisory journey.'),
    ('Insights', 'Reviewed financial-literacy archive and future articles.'),
    ('Book-a-call contact', 'Financial-goal and service enquiry.'),
    ('Pricing selector', 'Remove after approval; do not invent service tiers or prices.')
])

heading('3. Information architecture and routes')
heading('Primary navigation', 2)
for item in ['Home', 'About Us', 'Services', 'Online Services', 'Blog', 'Contact Us']:
    bullet(item)
heading('About Us dropdown', 2)
for item in ['Our Introduction', 'Awards', 'Certificates', 'Our Clients', 'Testimonials']:
    bullet(item)
heading('Services dropdown', 2)
for item in ['Financial Planning', 'Life Insurance', 'Health Insurance', 'Mutual Funds', 'Retirement Planning', 'Child Education Planning', 'Personal Accidental Policy', 'General Insurance', 'Employer Employee Insurance']:
    bullet(item)
heading('Online Services dropdown', 2)
for item in ['Newsletters', 'LIC Registered User', 'LIC Pay Premium Direct', 'Fundz Bazar', 'NJ E-Wealth Account', 'NJ Client Desk', 'Niva Bupa Renewal', 'Star Health Renewal', 'HDFC Ergo Renewal', 'Android App', 'iOS App']:
    bullet(item)
heading('Canonical routes', 2)
for item in ['/', '/about-us', '/about-us/awards', '/about-us/certificates', '/about-us/our-clients', '/about-us/testimonials', '/services', '/services/[service-slug]', '/online-services', '/newsletters', '/blog', '/blog/[article-slug]', '/contact-us', '/privacy-policy', '/terms-of-service', '/disclaimer']:
    bullet(item)
heading('Legacy compatibility', 2)
table([
    ('/about', '/about-us'),
    ('/contact', '/contact-us, retaining a valid service query'),
    ('/insights', '/blog'),
    ('/cases', '/about-us/our-clients'),
    ('/blog-post', '/blog'),
    ('Legacy Ghanchi root service/article slugs', 'Redirect only known records to their canonical routes.'),
    ('Fictional Kora case-detail slugs', 'Return 404 rather than imply they are Ghanchi clients.')
], ('Legacy destination', 'Canonical behavior'))

heading('4. Homepage: header')
heading('Visual and interaction behavior to keep', 2)
for item in ['White floating navigation pill.', 'Fixed positioning and backdrop blur.', 'Logo placement and spacing.', 'Dark right-side CTA with circular advisor image.', 'Current-page state.', 'Animated mobile drawer and Escape behavior.']:
    bullet(item)
heading('Ghanchi content mapping', 2)
for item in ['Kora logo becomes Ghanchi Investments.', 'Kora links become the six Ghanchi primary navigation items.', 'Book a call becomes Get in touch.', 'Header image becomes Chandrakant’s genuine image.', 'About Us, Services, and Online Services receive dropdowns.']:
    bullet(item)
p('Desktop dropdown design: the parent label remains a real page link and a separate disclosure control opens the submenu. This avoids hiding the Services and About landing pages behind hover-only behavior. Dropdowns support pointer and keyboard input, Escape closes them, and focus returns to the disclosure control.')
p('Mobile design: parent pages remain accessible, child lists expand inside the existing menu, long lists scroll within the viewport, selecting any destination closes the navigation, and Escape closes the entire menu and returns focus to the menu button.')

heading('5. Homepage: hero scene')
heading('Visual and motion behavior to keep', 2)
for item in ['Full-screen rounded media panel.', 'Initial media scale animation.', 'Pinned background while the visitor scrolls.', 'Scroll-linked background scale and blur.', 'Foreground vertical motion.', 'Word-by-word heading reveal.', 'Primary and secondary CTA positioning.', 'Lower-left trust area and lower-right featured card.']:
    bullet(item)
heading('Proposed hero content', 2)
p('Heading: “Your trusted partner for financial planning and protection.”')
p('Description: “Ghanchi Investments helps families plan, protect, and invest with confidence.”')
p('Primary CTA: “Our Services” → homepage Services section.')
p('Secondary CTA: “Get in touch” → Contact Us page.')
heading('Hero media decision', 2)
p('The abstract Kora-style visual may remain initially because it preserves the intended art direction and does not depict a fictitious customer or employee. Before launch, the owner may retain it, replace it with an approved Ghanchi image, commission professional imagery, or configure an approved video. Media replacement must not alter the hero animation.')

heading('6. Homepage: trust indicator and marquee')
heading('Trust indicator', 2)
for item in ['Retain overlapping circular badges, five mint dots, trust copy, and placement.', 'Use real client initials rather than Kora stock portraits or unrelated faces.', 'Show “Trusted by 1,200+ clients”.', 'Show “5.0/5 — Google rating on existing site” with a note that it is not a live feed.', 'Do not add a review count unless a verified source is connected.']:
    bullet(item)
heading('Marquee', 2)
p('Replace fictional Kora customer logos with verified client locations: India, UAE, and USA. These locations are supported by the current Ghanchi content and testimonials. Do not add more countries without owner confirmation or source evidence.')
p('Retain ticker duplication, mask fade, animation speed, typography, spacing, and reduced-motion fallback.')

heading('7. Homepage: featured newsletter card')
heading('Purpose', 2)
p('The existing Kora case-study card becomes a newsletter card without changing its shape, image/text split, badge, statistic row, hover behavior, or arrow treatment.')
heading('Initial content', 2)
for item in ['Badge: Latest in archive.', 'Title: November 2021 Newsletter.', 'Date: November 2021.', 'CTA: Read edition.', 'Destination: Newsletters page or the verified original newsletter URL.']:
    bullet(item)
p('November 2021 is currently the newest verified source issue. It must not be shown as a 2026 publication or as newly uploaded. Once Strapi is connected, the card can select the latest published available newsletter by issue month.')

heading('8. Homepage: statement and comparison')
heading('Statement scene', 2)
p('Retain the sticky scene, word-by-word interpolation, blur-to-sharp transition, responsive line-break handling, and reduced-motion presentation.')
p('Proposed statement: “Making goal-based financial advice accessible to all.” This comes from the existing Ghanchi vision.')
heading('Before/after comparison', 2)
table([
    ('Before heading', 'Financial decisions made in isolation.'),
    ('Before points', 'Goals planned but never implemented; insurance coverage that leaves gaps; investments scattered without a strategy; retirement approached without a clear plan.'),
    ('After heading', 'Every decision aligned with your goals.'),
    ('After points', 'Every goal backed by a clear plan; protection considered for the family; a portfolio aligned with risk profile; a retirement plan reviewed as needs change.')
], ('Element', 'Proposed content'))
p('Retain the card scale, overlap, horizontal movement, mint/dark contrast, check/cross icons, and sticky timeline. Avoid certainty language such as guaranteed wealth, guaranteed returns, or a corpus that lasts for life.')

heading('9. Homepage: services and animated chart')
heading('Section presentation', 2)
for item in ['Large Services heading.', 'Introductory paragraph with mint-highlighted phrase.', 'Animated two-bar chart.', 'Numbered split service cards.', 'Image hover and glass overlay.', 'Deliverable pills and full-width CTA.', 'Scroll reveal timing and responsive stacking.']:
    bullet(item)
p('Proposed introduction: “Your goals deserve more than isolated decisions. We bring planning, protection, and investing together around your life.”')
heading('Planning chart replacement', 2)
p('The revenue chart cannot remain a performance comparison because no equivalent Ghanchi performance data is verified. It becomes an illustration of the advisory journey while retaining the animation.')
table([
    ('First bar', 'Long-term goals — Retirement and education.'),
    ('Second bar', 'Today’s needs — Protection and cash flow.'),
    ('Axis labels', 'Understand, Plan, Implement, Review.'),
    ('Disclosure', 'An illustration of the planning journey — not a return forecast or performance comparison.')
], ('Chart element', 'Ghanchi mapping'))
heading('Nine service cards', 2)
p('Each service receives a title, short heading, sourced/conservative description, icon, illustrative or approved image, “What we discuss” list, service detail link, and enquiry CTA that preselects the service. The overlay uses Ghanchi’s vision or a clearly general testimonial rather than implying a quote relates to that specific product.')
for item in ['Financial Planning', 'Life Insurance', 'Health Insurance', 'Mutual Funds', 'Retirement Planning', 'Child Education Planning', 'Personal Accidental Policy', 'General Insurance', 'Employer Employee Insurance']:
    numbered(item)
p('The section becomes longer than Kora’s five-card version, but card design, spacing rhythm, responsive behavior, and interactions remain the same.')

heading('10. Homepage: advisory process')
table([
    ('Understand', 'Goals, objectives, current circumstances, and concerns.'),
    ('Assess', 'Risk appetite, income, expenses, liabilities, and cash flows.'),
    ('Plan', 'Customized recommendations around the client’s specific needs.'),
    ('Review', 'Implementation assistance and annual review.')
], ('Phase', 'Business meaning'))
heading('Interactions to keep', 2)
for item in ['Four-column spring accordion.', 'Numbering and progress dots.', 'Image/video presentation.', 'Circular play button.', 'Modal dialog.', 'Three operating-principle cards.']:
    bullet(item)
heading('Supporting principles', 2)
for item in ['Online access.', 'Annual review.', 'Personal service.']:
    bullet(item)
heading('Video fallback', 2)
p('Until an approved process video is supplied, keep the visual panel and play interaction. The dialog must state that the video is not connected, explain the four-step approach, and link to the genuine Ghanchi YouTube channel. It must not show a fictitious duration or silently pretend a video is available.')

heading('11. Homepage: advisor and awards')
heading('Advisor component', 2)
p('Retain the mint rounded section, numbered profile rows, circular image, plus icon, animated dialog, biography, and profile link. Initially show Chandrakant B. Ghanchi only, using a genuine source-site image and approved biography. Add staff only when real names, roles, biographies, and images are provided. Do not invent people to fill Kora’s original grid.')
heading('Former hiring panel', 2)
p('Retain the entire panel but change its meaning from recruitment to recognition.')
for item in ['Heading: Recognition. Built on service.', 'Image: genuine Ghanchi awards archive image.', 'Description: short statement about the business journey and service.', 'CTA: View Awards → /about-us/awards.']:
    bullet(item)

heading('12. Homepage: testimonials, rating, statistics, and founder CTA')
heading('Presentation to keep', 2)
for item in ['Large image-led feature panel.', 'Glass featured-quote card.', 'Large rating number and dots.', 'Expandable testimonial cards.', 'Animated statistics.', 'Founder callout and CTA.']:
    bullet(item)
heading('Approved testimonial sources', 2)
for item in ['Sumit Jain — L&T Infotech', 'Parvez Shaikh — East-west Freight Carriers Ltd', 'Deepak Salunkhe — Oberoi Realty', 'Ranbir Singh — USA', 'Neeta Agrawal — Syntel', 'Manju Rajvanshi — St. Xavier’s High School', 'Begum Dilshad — Home Maker', 'Ashu Rajvanshi — Acupressure Therapist', 'Firoz Shaikh — Firoz Dance Academy', 'Aaloak Singh Negi — TCS', 'Vikram Sawant — Dubai, UAE']:
    bullet(item)
p('Employer names are individual affiliations, not claims that the organizations are corporate clients or endorse Ghanchi Investments. Use initials when approved portraits are unavailable. Retain source URLs and owner approval metadata for every quote.')
heading('Statistics', 2)
for item in ['15+ years of experience.', '1,200+ clients served.', '12+ awards.', '5.0/5 rating reported on the existing website.']:
    bullet(item)
p('These are owner-selected editorial values and must remain editable. Do not create review counts or imply live synchronization.')

heading('13. Homepage: featured client-community panel')
p('The featured Kora case-study component remains, but fictional company results are replaced with truthful business context.')
table([
    ('Eyebrow', 'Serving clients since 2009.'),
    ('Heading', 'A personal approach, for clients in India and abroad.'),
    ('Facts', 'Focus: Your goals; Clients: Individuals and families; Review: Annually.'),
    ('Large statistics', '15+ years and 1,200+ clients.'),
    ('Image', 'Chandrakant or an approved Ghanchi business image.'),
    ('Service chips', 'Financial Planning, Insurance, Mutual Funds, and calculated remaining count.'),
    ('CTA', 'Meet Our Clients.')
], ('Component area', 'Proposed content'))
p('Preserve narrative/card layout, facts row, large-number styling, image hover, service chips, testimonial area if used, and CTA. Do not use fictional revenue, investment, or customer outcomes.')

heading('14. Pricing decision')
p('Remove the Kora pricing section because Ghanchi Investments does not publish equivalent packages, and inventing service tiers, premiums, fees, or monthly subscriptions would be misleading. This is the only currently approved section removal.')
p('After removal, adjust the surrounding rhythm so the FAQ follows naturally, there is no empty visual gap, and footer scroll calculations remain correct. No other section is removed without owner approval.')

heading('15. FAQ, blog, contact, and footer')
heading('FAQ', 2)
p('Retain the heading, category tabs, animated accordion, keyboard arrow behavior, and help/advisor card. Categories: General, Planning, Protection, Online Services. Topics include services, client types, office location, planning process, implementation and annual review, insurance considerations, claim assistance, external portals, static rating disclosure, and market risk.')
heading('Blog', 2)
p('Retain Kora’s image-led cards, category pills, dates, filters, article detail layout, and hover interactions. Publish only legitimate reviewed financial-literacy content. Initial verified archive examples are “Single Woman, Retiring Solo Is a Dream Retirement Life (but) With Proper Planning!” and “You Don’t Have to Be Rich to Retire Rich!” Preserve their 2021 dates, label condensed text as adaptations, retain source links, and do not attribute them to Chandrakant unless confirmed. Do not migrate unrelated casino/gambling posts or automatically consume the current WordPress feed.')
heading('Contact panel', 2)
for item in ['Keep the full rounded image background, large left heading, benefits list, testimonial, glass form, pill controls, statistics, and status messages.', 'Fields: Name, Email, Phone, Service interest, optional financial goal, enquiry/message, and explicit consent.', 'Remove annual revenue and growth-consulting fields.', 'Preserve client and server validation, honeypot, same-origin JSON checks, size limits, loading/error/success states, query preselection, and reduced-motion behavior.', 'Never show a successful delivery unless the configured provider confirms delivery; missing provider returns an honest unavailable state.']:
    bullet(item)
heading('Footer', 2)
for item in ['Keep the rounded cream card, mint circular reveal, scroll-linked scale, newsletter form, advisor contact, copy-email interaction, social icons, navigation hover, and large Ghanchi wordmark.', 'Use CBD Belapur address, both verified phone numbers, both verified email addresses, and genuine Facebook, Instagram, LinkedIn, and YouTube links.', 'Include draft legal pages and clear market-risk/insurance disclosures for owner and legal review.']:
    bullet(item)

heading('16. Inner pages')
table([
    ('About Us', 'Visual intro, Chandrakant section, history since 2009, 1,200+ clients, vision, mission, values, client categories, education content, and contact CTA.'),
    ('Awards', 'Responsive genuine-image gallery. Use neutral archive labels until titles, organizations, and dates are transcribed and approved.'),
    ('Certificates', 'Genuine certificate archive with historical-status notice. Do not imply current registration validity until reviewed.'),
    ('Our Clients', 'Client categories, India/UAE/USA, 1,200+ count, and explicit explanation that employer affiliations are not corporate-client claims.'),
    ('Testimonials', 'Full attributed quotes, initials where portraits are absent, source tracking, and no invented rating attached to individual quotes.'),
    ('Services listing', 'All nine services in the Kora inner-card design.'),
    ('Service detail', 'Title, description, hero image, discussion points, general testimonial with disclosure, regulatory/product warning, related services, and preselected enquiry CTA.'),
    ('Online Services', 'Provider grid, external-link indicator, security notice, apps, renewals, account links, and newsletter access. No embedded provider login.'),
    ('Newsletters', 'Chronological source archive. November 2021 is newest verified. Malformed November 2020 and June 2020 links display as unavailable until corrected by owner.'),
    ('Blog', 'Reviewed categories and articles, original dates, adaptation notice, source link, empty states, and financial disclaimer.'),
    ('Contact Us', 'Full form, verified address, both phones, both emails, social links, and optional approved map later.'),
    ('Legal pages', 'Privacy, terms, and disclaimer drafts; noindex until approved for launch.')
], ('Page', 'Content and behavior'))

heading('17. Content truth and approval policy')
heading('Owner-approved or source-supported content', 2)
for item in ['Business name.', 'Founded/incorporated in 2009.', 'Chandrakant B. Ghanchi.', '1,200+ clients.', 'Nine listed services.', 'Vision, mission, and advisory approach.', 'Contact information and social links.', 'Actual testimonials.', 'Awards and certificate images.', 'Online-service links.', 'Newsletter archive.', 'Legitimate historical financial articles.', 'India, UAE, and USA client presence.']:
    bullet(item)
heading('Content that requires qualification or pre-launch review', 2)
for item in ['15+ years.', '12+ awards.', '5.0 Google rating.', 'Current regulatory titles or designations.', 'Product, tax, policy-benefit, or return wording.', 'Certificate validity and expiry.', 'External provider destinations.', 'Image reuse rights.', 'Legal and privacy disclosures.']:
    bullet(item)
heading('Content prohibited without new approval/evidence', 2)
for item in ['Kora revenue claims, prices, clients, team, and outcomes.', 'Guaranteed investment returns or financial security.', 'Invented countries, staff, testimonials, awards, certificates, review counts, or customers.', 'Corporate-client claims inferred from an individual’s employer.', 'Casino/gambling content.', 'Newer newsletters or repaired URLs guessed from patterns.']:
    bullet(item)

heading('18. Future Strapi management')
p('The frontend should be prepared for Strapi, but infrastructure must not be provisioned until separately approved. Strapi will manage editorial data; code will continue to define safe layouts, animations, validation, and integrations.')
heading('Single types', 2)
for item in ['Site Settings', 'Navigation', 'Homepage', 'Contact Settings', 'Footer Settings']:
    bullet(item)
heading('Collection types', 2)
for item in ['Page', 'Service', 'Advisor', 'Testimonial', 'Award', 'Certificate', 'Client Reference', 'Client Location', 'Article', 'Category', 'Newsletter', 'Online Service', 'Statistic', 'FAQ', 'Legal Page']:
    bullet(item)
heading('Editable without code changes', 2)
for item in ['Site name and approved logo media.', 'Header/footer labels, menu items, child links, order, and enabled state.', 'Hero heading, description, CTAs, media, overlay, trust content, locations, rating note, and featured newsletter.', 'Comparison copy, service records, service images, benefits, and order.', 'Process phases and approved icon/motion presets.', 'Advisor profiles, awards, certificates, testimonials, statistics, FAQs, articles, newsletters, and online links.', 'Section ordering and enabled state within an allowlisted registry.', 'Contact details, social links, SEO metadata, legal pages, and disclosures.']:
    bullet(item)
heading('Never editable as unrestricted content', 2)
for item in ['Arbitrary JavaScript or executable HTML.', 'Raw CSS or unbounded motion code.', 'Database queries.', 'API tokens, webhook URLs, or secrets.', 'Unknown section/component identifiers.', 'Unsafe URL schemes or arbitrary iframe hosts.', 'Required security validation, consent, or honeypot behavior.']:
    bullet(item)
p('Use a typed allowlisted section registry rather than a freeform page builder. New layout types and executable behavior continue to require development. The hero, statement, scroll spacer, and comparison remain one indivisible intro scene so editors cannot break the motion timeline.')

heading('19. Implementation sequence after approval')
steps = [
    'Record the owner’s changes to this document and mark the final section mappings approved.',
    'Compare the isolated Ghanchi worktree with the clean Kora backup and preserve the backup unchanged.',
    'Finalize source-backed content, remove any draft assumptions, and confirm imagery/rights.',
    'Complete header, footer, hero, marquee, newsletter card, statement, and comparison while preserving motion code.',
    'Complete all nine service cards and service detail pages.',
    'Complete advisory process, advisor dialog, awards panel, testimonials, statistics, and client-community panel.',
    'Remove pricing only and verify page rhythm and scroll timelines.',
    'Complete About subpages, Online Services, Newsletters, Blog, Contact, legal pages, metadata, sitemap, robots, redirects, and 404 behavior.',
    'Complete form and API validation with mocked delivery and honest unconfigured states.',
    'Update automated tests and the separate Strapi implementation handoff.',
    'Run content validation, TypeScript, production build, full browser tests, responsive visual review, keyboard review, reduced-motion review, and motion checkpoints.',
    'Present the isolated local preview for owner review and apply requested refinements.',
    'Do not commit, push, deploy, provision Strapi, or change DNS without separate authorization.'
]
for step in steps:
    numbered(step)

heading('20. Acceptance criteria')
for item in [
    'No Kora-specific business names, emails, prices, claims, fictional customers, or team members remain in published Ghanchi content.',
    'No unrelated gambling/spam material is migrated.',
    'All major Kora visual components remain except explicitly approved pricing.',
    'Hero pinning, statement reveal, comparison motion, accordions, dialogs, testimonial expansion, counters, image hovers, and footer motion retain their behavior.',
    'Every published route has one H1, a working skip target, Ghanchi metadata, no broken images, and no horizontal overflow.',
    'Navigation works with mouse, keyboard, touch, and Escape at desktop and mobile widths.',
    'Forms remain safe before hydration, validate on both client and server, preserve values after failure, and never falsely report delivery.',
    'Reduced-motion visitors can read and use all content without hidden or pinned-only states.',
    'Visual checks cover 1440px, 1200px, 810px, and 390px after scrolling to activate reveal states.',
    'Owner reviews wording, statistics, gallery captions, external links, legal disclosures, and imagery before launch.',
    'Strapi is not described as connected until server-side adapters, preview, publication, and revalidation are actually implemented and tested.'
]: bullet(item)

heading('21. Decisions requested from the owner')
questions = [
    'Approve or revise the hero heading: “Your trusted partner for financial planning and protection.”',
    'Keep the abstract hero artwork initially, or provide/select an approved Ghanchi replacement?',
    'Approve India, UAE, and USA as the only marquee locations until more are confirmed?',
    'Approve November 2021 as the honestly labelled latest item in the current newsletter archive?',
    'Approve repurposing the Kora chart as a planning-journey illustration?',
    'Approve displaying only Chandrakant until verified staff information is provided?',
    'Approve converting the hiring panel into awards and recognition?',
    'Approve converting the featured case-study panel into client-community and credibility content?',
    'Confirm complete pricing-section removal?',
    'Approve static 5.0 rating wording until a permitted live integration exists?',
    'Approve neutral gallery captions until award/certificate details are transcribed?',
    'Approve publishing only reviewed legitimate archive articles with adaptation labels?'
]
for question in questions:
    numbered(question)

heading('22. Approval record')
p('Owner response / requested changes:')
for _ in range(8):
    p('________________________________________________________________________________')
p('Approved by: __________________________________    Date: ________________________')
p('Implementation may resume after the above decisions are recorded and approved.')

doc.save(OUTPUT)
with ZipFile(OUTPUT) as archive:
    ElementTree.fromstring(archive.read('word/document.xml'))
loaded = Document(OUTPUT)
text = '\n'.join(paragraph.text for paragraph in loaded.paragraphs)
for required in ['Transformation objective', 'Homepage: hero scene', 'Pricing decision', 'Future Strapi management', 'Decisions requested from the owner']:
    if required.lower() not in text.lower():
        raise RuntimeError(f'Missing required section: {required}')
print(f'Generated and validated {OUTPUT.name}: {len(loaded.paragraphs)} paragraphs, {len(loaded.tables)} tables.')
